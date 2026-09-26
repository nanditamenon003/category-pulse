"""
A store's own data: the Excel template, the sample file, and reading uploads.

The template has one sheet per table. Targets and Sales are required; Stock,
Visitors, Loyalty and Settings are optional, and each one unlocks more of
the app (see feature_checklist). Column names are matched loosely (case,
spaces, and anything in brackets are ignored, and common alternatives like
"Qty" for "Units" are accepted), so an export from the store's own system
needs little tidying. CSV files work too, one per table, named after it
(e.g. sales.csv).

The sample file is the simulated demo month written into the template, so
anyone can see exactly how to fill it in, or try the upload in one click.
"""

import functools
import hashlib
import io
import re
from datetime import date, datetime, time

import pandas as pd

from store import WEEKDAY_NAMES, StoreDataError, build_store, demo_store

# --- The template ---------------------------------------------------------------------------

# Sheet -> its column headers, as they appear in the template.
TEMPLATE = {
    "Targets": ["Line", "Category", "Floor (optional)", "Monthly target (units)",
                "Last year (units, optional)", "Sizes (optional)", "Core sizes (optional)"],
    "Sales": ["Date", "Hour (optional)", "Line", "Category", "Size (optional)", "Units",
              "Value (₹, optional)", "Bills (optional)"],
    "Stock": ["Date", "Hour (optional)", "Line", "Category", "Size (optional)", "Units on hand"],
    "Visitors": ["Date", "Hour (optional)", "Floor (optional)", "Visitors"],
    "Loyalty": ["Tier", "Line", "Category", "Share of bills (%)", "Average bill value (₹)",
                "Units per bill", "Takes up cross-sells (%)", "Preferred offer"],
    "Settings": ["Setting", "Value"],
}
SETTINGS_ROWS = ["Store name", "Delivery day", "Sale days"]

READ_ME = [
    ("Category Pulse: store data template", ""),
    ("", ""),
    ("How to fill it in", "One month of data. Keep each sheet's header row as it is. Targets and Sales "
                          "are required; the other sheets are optional and each one unlocks more of "
                          "the app. Leave a sheet empty if you don't have that data."),
    ("", ""),
    ("Targets", "One row per category: its line, the category (e.g. Polo), the floor it's on "
                "(e.g. Menswear), and its monthly target in units. Optional: last year's units; its "
                "sizes in shelf order, separated by commas (S, M, L, XL); its core sizes, the ones "
                "most shoppers need (M, L). If core sizes are left blank they're worked out from "
                "sales."),
    ("Sales", "One row per sale line: date, line, category and units. Optional: the hour (10 means "
              "10:00-11:00), the size, the value in rupees, and the number of bills. Rows can be "
              "per bill or already totalled; they're added up either way. If bills are only known "
              "per category and hour, put the count on one row and leave the others blank."),
    ("Stock", "Units on hand by category and size, counted at the close of a day (or at an hour, "
              "if you count hourly). One count on the last day is enough; a count every day also "
              "shows deliveries and what's usual."),
    ("Visitors", "Visitors counted per floor, by day or by hour. Needs Bills in the Sales sheet to "
                 "tell fewer visitors apart from fewer buyers. By hour, it also gives tomorrow's "
                 "busy hours."),
    ("Loyalty", "Loyalty tier figures per category (tier level only, never individual customers). "
                "Preferred offer is one of: outfit bundle, first-purchase discount, double points, "
                "early access, free alteration."),
    ("Settings", "Store name; the weekday deliveries arrive (e.g. Monday), used to spot missed "
                 "deliveries; and any sale days, as day numbers separated by commas (e.g. 13, 27)."),
    ("", ""),
    ("Dates", "Any normal date format works; day first (24/05/2026) is read as the day."),
    ("Privacy", "Only upload real company figures with your manager's approval. The app reads the "
                "file for your session only and doesn't save it, and the AI chat is switched off "
                "for uploaded data."),
]

# --- Reading uploads ------------------------------------------------------------------------

SHEETS = {"targets": "targets", "target": "targets", "sales": "sales", "stock": "stock",
          "inventory": "stock", "visitors": "visitors", "footfall": "visitors",
          "loyalty": "loyalty", "settings": "settings"}

_PRODUCT = {"category": "product", "product": "product", "product type": "product"}
_FLOOR = {"floor": "department", "department": "department", "zone": "department"}
COLUMNS = {
    "targets": {"line": "line", **_PRODUCT, **_FLOOR, "monthly target": "target", "target": "target",
                "last year": "last_year", "ly": "last_year", "sizes": "sizes", "core sizes": "core_sizes"},
    "sales": {"date": "date", "bill date": "date", "hour": "hour", "time": "hour", "line": "line",
              **_PRODUCT, "size": "size", "units": "units", "qty": "units", "quantity": "units",
              "units sold": "units", "value": "value", "sales value": "value", "net value": "value",
              "amount": "value", "net sales": "value", "bills": "transactions",
              "transactions": "transactions", "invoices": "transactions", "no of bills": "transactions"},
    "stock": {"date": "date", "hour": "hour", "time": "hour", "line": "line", **_PRODUCT,
              "size": "size", "units on hand": "units", "units": "units", "qty": "units",
              "stock": "units", "soh": "units", "closing stock": "units"},
    "visitors": {"date": "date", "hour": "hour", "time": "hour", **_FLOOR, "visitors": "visitors",
                 "footfall": "visitors", "walk-ins": "visitors", "walk ins": "visitors"},
    "loyalty": {"tier": "tier", "line": "line", **_PRODUCT,
                "share of bills": "share_of_transactions", "average bill value": "avg_basket_value",
                "units per bill": "avg_upt", "takes up cross-sells": "cross_sell_response_rate",
                "preferred offer": "preferred_offer_type"},
}

OFFERS = {"bundle": "bundle", "outfit bundle": "bundle",
          "first-purchase discount": "percentage_discount", "discount": "percentage_discount",
          "percentage discount": "percentage_discount",
          "double points": "loyalty_points_multiplier", "points": "loyalty_points_multiplier",
          "loyalty points multiplier": "loyalty_points_multiplier",
          "early access": "early_access", "free alteration": "complimentary_service",
          "alteration": "complimentary_service", "complimentary service": "complimentary_service"}
OFFER_WORDS = {"bundle": "outfit bundle", "percentage_discount": "first-purchase discount",
               "loyalty_points_multiplier": "double points", "early_access": "early access",
               "complimentary_service": "free alteration"}


def _plain(text):
    """'Value (₹, optional)' -> 'value'; 'Units_Sold' -> 'units sold'."""
    text = re.sub(r"\(.*?\)", "", str(text)).replace("_", " ").lower()
    return re.sub(r"\s+", " ", text).strip()


def _hour(value):
    """10, 10.0, '10', '10:00', '10:00-11:00' or a time -> 10. Blank stays blank."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (time, datetime)):
        return value.hour
    match = re.match(r"\s*(\d{1,2})", str(value))
    return int(match.group(1)) if match else value  # unreadable: build_store says so


def _read_files(files):
    """[(file name, bytes)] -> {table name: DataFrame}, plus notes about what was skipped."""
    tables, notes = {}, []
    for name, data in files:
        lower = name.lower()
        if lower.endswith((".xlsx", ".xlsm")):
            try:
                sheets = pd.read_excel(io.BytesIO(data), sheet_name=None, engine="openpyxl")
            except Exception as e:
                raise StoreDataError(f"{name} couldn't be opened as an Excel file ({e}).")
        elif lower.endswith(".csv"):
            try:
                sheets = {re.sub(r"\.csv$", "", name, flags=re.I): pd.read_csv(io.BytesIO(data))}
            except Exception as e:
                raise StoreDataError(f"{name} couldn't be read as a CSV file ({e}).")
        else:
            raise StoreDataError(f"{name} isn't an Excel (.xlsx) or CSV file.")
        for sheet, df in sheets.items():
            table = SHEETS.get(_plain(sheet))
            if table is None:
                if _plain(sheet) not in ("read me", "readme", "instructions"):
                    notes.append(f"Skipped the sheet \"{sheet}\" (not one of the template's sheets).")
                continue
            df = df.dropna(how="all")
            if not df.empty:
                tables[table] = df
    return tables, notes


def _rename(df, table, notes):
    mapping, ignored = {}, []
    for column in df.columns:
        canonical = COLUMNS[table].get(_plain(column))
        if canonical and canonical not in mapping.values():
            mapping[column] = canonical
        else:
            ignored.append(str(column))
    if ignored and not all(c.startswith("Unnamed") for c in ignored):
        notes.append(f"{table.title()}: ignored the column(s) {', '.join(ignored)}.")
    df = df[list(mapping)].rename(columns=mapping)
    if "hour" in df.columns:
        df["hour"] = df["hour"].map(_hour)
    return df


def _settings(df):
    """The Settings sheet -> (store name, delivery weekday or None, sale days)."""
    values = {}
    if df is not None and df.shape[1] >= 2:
        for _, row in df.iterrows():
            values[_plain(row.iloc[0])] = row.iloc[1]
    name = values.get("store name")
    name = str(name).strip() if pd.notna(name) and str(name).strip() else "Your store"

    delivery = values.get("delivery day")
    weekday = None
    if pd.notna(delivery) and str(delivery).strip():
        wanted = str(delivery).strip().lower()[:3]
        matches = [i for i, day in enumerate(WEEKDAY_NAMES) if day.lower().startswith(wanted)]
        if not matches:
            raise StoreDataError(f"Settings: the delivery day {delivery!r} isn't a weekday name.")
        weekday = matches[0]

    sale_days, raw = [], values.get("sale days")
    if isinstance(raw, (datetime, date)):
        sale_days = [raw.day]
    elif pd.notna(raw) and str(raw).strip():
        for part in (p.strip() for p in re.split(r"[,;]", str(raw)) if p.strip()):
            if re.fullmatch(r"\d{1,2}(\.0)?", part):  # a day number, e.g. 13
                sale_days.append(int(float(part)))
                continue
            when = pd.to_datetime(part, errors="coerce", dayfirst=True)  # a date, e.g. 13/05/2026
            if pd.isna(when):
                raise StoreDataError(f"Settings: the sale day {part!r} isn't a day number or a date.")
            sale_days.append(when.day)
    return name, weekday, sale_days


def read_upload(files):
    """
    Reads uploaded files ([(name, bytes)]: one Excel template, or CSV files
    named after the tables) into a Store. Returns (store, notes). Raises
    StoreDataError with a plain explanation if the data can't be used.
    """
    tables, notes = _read_files(files)
    missing = [t.title() for t in ("targets", "sales") if t not in tables]
    if missing:
        raise StoreDataError(f"No {' or '.join(missing)} found. The {' and '.join(missing)} "
                             f"sheet{'s need' if len(missing) > 1 else ' needs'} at least one row.")
    renamed = {t: _rename(df, t, notes) for t, df in tables.items() if t != "settings"}
    name, delivery_weekday, sale_days = _settings(tables.get("settings"))

    loyalty = renamed.get("loyalty")
    if loyalty is not None and "preferred_offer_type" in loyalty.columns:
        loyalty["preferred_offer_type"] = loyalty["preferred_offer_type"].map(
            lambda v: OFFERS.get(_plain(v), _plain(v).replace(" ", "_")))

    fingerprint = hashlib.sha1(b"".join(data for _, data in files)).hexdigest()[:12]
    store = build_store(
        renamed["targets"], renamed["sales"], renamed.get("stock"), renamed.get("visitors"), loyalty,
        name=name, store_id=f"upload-{fingerprint}", sale_days=sale_days,
        delivery_weekday=delivery_weekday,
    )
    return store, notes


# --- What the data unlocks ----------------------------------------------------------------------

def feature_checklist(store):
    """
    Each feature, whether this store's data supports it, and (when it's off)
    what would switch it on.
    """
    has_bills_and_visitors = store.has_footfall and store.has_transactions
    features = [
        ("Month-to-date pace and status for every category", True, ""),
        ("Contribution report and end-of-day summary", True, ""),
        ("Rupee values in the contribution report", store.has_value, "add Value to Sales"),
        ("Step through the day hour by hour", store.hourly, "add Hour to Sales"),
        ("Broken size runs and core sizes", store.has_sizes and store.has_stock,
         "add Size to Sales and Stock" if store.has_stock else "add a Stock sheet with sizes"),
        ("Stock problems and last-piece alerts", store.has_stock, "add a Stock sheet"),
        ("Deliveries and usual stock levels", store.has_stock_history,
         "count stock on more than one day"),
        ("Missed deliveries and run-out warnings", store.has_stock_history and store.delivery_weekday is not None,
         "add a Delivery day in Settings" if store.has_stock_history else "daily stock counts and a Delivery day in Settings"),
        ("Fewer visitors or fewer buyers?", has_bills_and_visitors,
         "add a Visitors sheet" + ("" if store.has_transactions else " and Bills to Sales")),
        ("Tomorrow's busy hours and floor split", store.has_visitor_hours, "add visitors by hour"),
        ("Loyalty tier targeting at the till", store.has_loyalty, "add a Loyalty sheet"),
    ]
    return [(name, on, "" if on else hint) for name, on, hint in features]


# --- Writing the template and the sample ------------------------------------------------------------

def _workbook(frames):
    """{sheet: DataFrame} -> .xlsx bytes, with a Read me sheet, bold headers and sensible widths."""
    from openpyxl.styles import Alignment, Font

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pd.DataFrame(READ_ME).to_excel(writer, sheet_name="Read me", index=False, header=False)
        sheet = writer.sheets["Read me"]
        sheet.column_dimensions["A"].width = 24
        sheet.column_dimensions["B"].width = 110
        for row in sheet.iter_rows():
            row[0].font = Font(bold=True)
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
        sheet["A1"].font = Font(bold=True, size=14)

        for name, df in frames.items():
            df.to_excel(writer, sheet_name=name, index=False)
            sheet = writer.sheets[name]
            sheet.freeze_panes = "A2"
            for i, column in enumerate(df.columns):
                cell = sheet.cell(row=1, column=i + 1)
                cell.font = Font(bold=True)
                letter = cell.column_letter
                sheet.column_dimensions[letter].width = max(12, len(str(column)) + 3)
                if pd.api.types.is_datetime64_any_dtype(df[column]) or column == "Date":
                    for row in range(2, len(df) + 2):
                        sheet.cell(row=row, column=i + 1).number_format = "DD/MM/YYYY"
    return buffer.getvalue()


@functools.lru_cache(maxsize=1)
def template_bytes():
    """The empty template: every sheet with its header row, plus the Read me."""
    frames = {name: pd.DataFrame(columns=columns) for name, columns in TEMPLATE.items()}
    frames["Settings"] = pd.DataFrame({"Setting": SETTINGS_ROWS, "Value": [None] * len(SETTINGS_ROWS)})
    return _workbook(frames)


@functools.lru_cache(maxsize=1)
def sample_bytes():
    """
    The simulated demo month, filled into the template: hourly sales by
    size, stock counted at each day's close, visitors by hour, loyalty
    tiers and settings. Uploading it shows the demo store as at the close
    of its latest day.
    """
    s = demo_store()
    t = TEMPLATE

    targets = pd.DataFrame([{
        t["Targets"][0]: s.category_line[c], t["Targets"][1]: s.category_product[c],
        t["Targets"][2]: s.category_department[c], t["Targets"][3]: s.targets[c],
        t["Targets"][4]: s.last_year.get(c), t["Targets"][5]: ", ".join(s.sizes[c]),
        t["Targets"][6]: ", ".join(s.core_sizes[c]),
    } for c in s.categories])

    # One row per category, hour and size that sold; each category-hour's
    # bills sit on its first row, as the Read me describes.
    sales = s.sales[s.sales["units_sold"] > 0].drop(columns="transactions", errors="ignore")
    sales = sales.merge(s.transactions, on=["day", "hour", "category"])
    sales = sales.sort_values(["day", "hour", "category"], kind="stable")
    first_row = ~sales.duplicated(["day", "hour", "category"])
    sales = pd.DataFrame({
        "Date": sales["day"].map(s.date), "Hour (optional)": sales["hour"],
        "Line": sales["line"], "Category": sales["category"].map(s.category_product),
        "Size (optional)": sales["size"], "Units": sales["units_sold"],
        "Value (₹, optional)": sales["value"].round(2),
        "Bills (optional)": sales["transactions"].where(first_row),
    })

    closing = s.stock[s.stock["hour"] == s.hours[-1]]
    stock = pd.DataFrame({
        "Date": closing["day"].map(s.date), "Line": closing["line"],
        "Category": closing["category"].map(s.category_product), "Size (optional)": closing["size"],
        "Units on hand": closing["units_remaining"],
    })

    visitors = pd.DataFrame({
        "Date": s.footfall["day"].map(s.date), "Hour (optional)": s.footfall["hour"],
        "Floor (optional)": s.footfall["zone"], "Visitors": s.footfall["visitors"],
    })

    loyalty = pd.DataFrame({
        "Tier": s.loyalty["tier"], "Line": s.loyalty["category"].map(s.category_line),
        "Category": s.loyalty["category"].map(s.category_product),
        "Share of bills (%)": s.loyalty["share_of_transactions"],
        "Average bill value (₹)": s.loyalty["avg_basket_value"],
        "Units per bill": s.loyalty["avg_upt"],
        "Takes up cross-sells (%)": s.loyalty["cross_sell_response_rate"],
        "Preferred offer": s.loyalty["preferred_offer_type"].map(OFFER_WORDS),
    })

    settings = pd.DataFrame({"Setting": SETTINGS_ROWS, "Value": [
        "Sample Store (from the sample file)", WEEKDAY_NAMES[s.delivery_weekday],
        ", ".join(str(d) for d in sorted(s.sale_days))]})

    return _workbook({"Targets": targets, "Sales": sales, "Stock": stock, "Visitors": visitors,
                      "Loyalty": loyalty, "Settings": settings})
